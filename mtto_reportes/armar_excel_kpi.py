"""
Arma un Excel sin encabezados a partir de un CSV descargado de la app.

Una fila por técnico. Solo la primera fila del reporte lleva el resto
de columnas; las siguientes solo traen el nombre en la columna O.
El historial se parte por "→" desde la columna AH, con el texto completo.

Uso:
    python armar_excel_kpi.py
    python armar_excel_kpi.py MANT_2026-10-05.csv
"""
import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook

CARPETA = Path(__file__).resolve().parent

# Columna de Excel (1 = A) -> nombre en el CSV
COLUMNAS = {
    2: "FECHA",          # B
    6: "AREA",           # F
    7: "MAQUINA",         # G
    8: "SOLICITANTE",    # H
    9: "HORA SOLICITA",  # I
    11: "PRIORIDAD",     # K
    12: "FALLA",         # L
    13: "TIPO",          # M
    16: "FECHA ENTREGA", # P
    17: "SOLICITANTE",   # Q
    19: "HORA DE ENTREGA",  # S
    27: "DURACION TOTAL (h)",  # AA
    28: "PROBLEMA",      # AB
    29: "MANTENIMIENTO", # AC
    30: "REFACCIONES",   # AD
}
COL_TECNICO = 15          # O
COL_HORAS = 21            # U
COL_MINUTOS = 22          # V
COL_HISTORIAL = 34        # AH


def csv_mas_reciente():
    archivos = list(CARPETA.glob("MANT_*.csv"))
    if not archivos:
        raise SystemExit("No hay archivos MANT_*.csv en " + str(CARPETA))
    return max(archivos, key=lambda p: p.stat().st_mtime)


def texto(valor):
    if valor is None or (isinstance(valor, float) and pd.isna(valor)):
        return ""
    s = str(valor).strip()
    if s.lower() == "nan":
        return ""
    return s


def horas_y_minutos(valor):
    s = texto(valor)
    if not s:
        return "", ""
    try:
        minutos = int(round(float(s.replace(",", "."))))
    except ValueError:
        return "", ""
    if minutos < 0:
        return "", ""
    return minutos // 60, minutos % 60


def tecnicos_de(valor):
    s = texto(valor)
    if not s:
        return [""]
    partes = [p.strip() for p in s.split("/") if p.strip()]
    return partes or [""]


def eventos_historial(valor):
    s = texto(valor)
    if not s:
        return []
    return [p.strip() for p in s.split("→") if p.strip()]


def escribir_filas(df, destino):
    wb = Workbook()
    ws = wb.active
    ws.title = "KPI"
    excel_fila = 1
    for _, reg in df.iterrows():
        horas, minutos = horas_y_minutos(reg.get("TIEMPO REPARACION (MIN)", ""))
        eventos = eventos_historial(reg.get("HISTORIAL", ""))
        for i_tec, tecnico in enumerate(tecnicos_de(reg.get("TECNICOS", ""))):
            if i_tec == 0:
                for col, nombre in COLUMNAS.items():
                    valor = texto(reg.get(nombre, ""))
                    if valor != "":
                        ws.cell(row=excel_fila, column=col, value=valor)
                if horas != "":
                    ws.cell(row=excel_fila, column=COL_HORAS, value=horas)
                    ws.cell(row=excel_fila, column=COL_MINUTOS, value=minutos)
                for i, evento in enumerate(eventos):
                    ws.cell(row=excel_fila, column=COL_HISTORIAL + i, value=evento)
            if tecnico != "":
                ws.cell(row=excel_fila, column=COL_TECNICO, value=tecnico)
            excel_fila += 1
    wb.save(destino)
    return excel_fila - 1


def main():
    if len(sys.argv) > 1:
        origen = Path(sys.argv[1])
        if not origen.is_absolute():
            origen = CARPETA / origen
    else:
        origen = csv_mas_reciente()
    if not origen.exists():
        raise SystemExit("No existe: " + str(origen))

    df = pd.read_csv(origen, encoding="utf-8-sig", dtype=str, keep_default_na=False)
    fecha = origen.stem.replace("MANT_", "")
    destino = origen.with_name("KPI_" + fecha + ".xlsx")
    n = escribir_filas(df, destino)
    print(f"{origen.name}: {len(df)} servicios -> {n} filas en {destino.name}")


if __name__ == "__main__":
    main()
